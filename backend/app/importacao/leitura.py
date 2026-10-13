import csv
import io
import zipfile
from dataclasses import dataclass

from openpyxl import load_workbook

COLUNAS = (
    "cpf",
    "nome_cliente",
    "data_nascimento",
    "sexo",
    "estado_civil",
    "email",
    "telefone",
    "cep_cliente",
    "logradouro",
    "numero",
    "complemento",
    "bairro",
    "cidade_cliente",
    "uf_cliente",
    "cnpj",
    "razao_social",
    "nome_fantasia",
    "setor",
    "cep_empresa",
    "cidade_empresa",
    "uf_empresa",
    "matricula",
    "cargo",
    "data_admissao",
    "salario_bruto",
    "margem_disponivel",
)


class ArquivoInvalido(ValueError):
    pass


@dataclass(frozen=True)
class LinhaBruta:
    numero: int
    valores: dict[str, object]


def _normalizar_cabecalho(celula: object) -> str:
    return str(celula or "").strip().lower().replace(" ", "_")


def _montar(linhas: list[tuple[int, list[object]]]) -> list[LinhaBruta]:
    if not linhas:
        raise ArquivoInvalido("O arquivo está vazio.")
    numero_cab, cabecalho_bruto = linhas[0]
    cabecalho = [_normalizar_cabecalho(c) for c in cabecalho_bruto]
    faltando = [c for c in COLUNAS if c not in cabecalho]
    if faltando:
        raise ArquivoInvalido(
            f"O cabeçalho (linha {numero_cab}) não tem as colunas obrigatórias: {', '.join(faltando)}."
        )
    indices = {c: cabecalho.index(c) for c in COLUNAS}
    saida = []
    for numero, celulas in linhas[1:]:
        if all(c is None or str(c).strip() == "" for c in celulas):
            continue
        valores = {c: (celulas[i] if i < len(celulas) else None) for c, i in indices.items()}
        saida.append(LinhaBruta(numero, valores))
    return saida


def _ler_xlsx(conteudo: bytes) -> list[LinhaBruta]:
    try:
        livro = load_workbook(io.BytesIO(conteudo), read_only=True, data_only=True)
    except (zipfile.BadZipFile, KeyError, OSError) as e:
        raise ArquivoInvalido("Não foi possível abrir o .xlsx. O arquivo pode estar corrompido.") from e
    try:
        planilha = livro.worksheets[0]
        linhas: list[tuple[int, list[object]]] = []
        for numero, celulas in enumerate(planilha.iter_rows(values_only=True), start=1):
            if not linhas and all(c is None for c in celulas):
                continue
            linhas.append((numero, list(celulas)))
        return _montar(linhas)
    finally:
        livro.close()


def _decodificar(conteudo: bytes) -> str:
    for codificacao in ("utf-8-sig", "cp1252"):
        try:
            return conteudo.decode(codificacao)
        except UnicodeDecodeError:
            continue
    raise ArquivoInvalido("Não foi possível ler o texto do CSV. Salve o arquivo em UTF-8.")


def _ler_csv(conteudo: bytes) -> list[LinhaBruta]:
    texto = _decodificar(conteudo)
    primeira = texto.split("\n", 1)[0]
    separador = ";" if primeira.count(";") >= primeira.count(",") else ","
    leitor = csv.reader(io.StringIO(texto), delimiter=separador)
    linhas = [(leitor.line_num, list(celulas)) for celulas in leitor]
    return _montar([(n, c) for n, c in linhas if c or n == 1])


def ler_planilha(nome_arquivo: str, conteudo: bytes) -> list[LinhaBruta]:
    nome = nome_arquivo.lower()
    if nome.endswith(".xlsx"):
        return _ler_xlsx(conteudo)
    if nome.endswith(".csv"):
        return _ler_csv(conteudo)
    raise ArquivoInvalido("Formato não suportado. Envie .xlsx ou .csv.")
